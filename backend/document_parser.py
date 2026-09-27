from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file type. Upload a TXT, MD, PDF, or DOCX file.")

    if extension in {".txt", ".md"}:
        return content.decode("utf-8-sig").strip()

    if extension == ".pdf":
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    table_text = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + table_text).strip()